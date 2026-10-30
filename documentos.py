"""Lectura de documentos públicos sin ejecutar adjuntos ni extraer rutas de ZIP."""
from io import BytesIO
from pathlib import Path
import os
import re
import subprocess
import tempfile
import zipfile
from xml.etree import ElementTree


def docx(datos):
    with zipfile.ZipFile(BytesIO(datos)) as z:
        i=z.getinfo('word/document.xml')
        if i.file_size>32_000_000:raise ValueError('DOCX excede tamaño seguro')
        root=ElementTree.fromstring(z.read(i))
        return '\n'.join(' '.join(t.text or '' for t in p.iter() if t.tag.endswith('}t'))
                         for p in root.iter() if p.tag.endswith('}p'))


def pdf(datos,ocr=False,progreso=None):
    if not datos.startswith(b'%PDF-'):raise ValueError('Respuesta no es PDF')
    r=subprocess.run(['pdftotext','-layout','-','-'],input=datos,capture_output=True,check=True,timeout=35)
    texto=r.stdout.decode('utf-8')
    if texto.strip():return texto
    if not ocr:raise ValueError('PDF sin texto: requiere OCR; no se declara vacío')
    with tempfile.TemporaryDirectory() as d:
        source=Path(d)/'documento.pdf';source.write_bytes(datos)
        info=subprocess.run(['pdfinfo',str(source)],capture_output=True,text=True,check=True,timeout=5)
        m=re.search(r'^Pages:\s*(\d+)',info.stdout,re.M)
        if not m:raise ValueError('No se pudo obtener cantidad de páginas para OCR')
        paginas=[]
        for pagina in range(1,int(m[1])+1):
            prefix=Path(d)/'pagina'
            subprocess.run(['pdftoppm','-f',str(pagina),'-l',str(pagina),'-scale-to','1600','-singlefile','-png',str(source),str(prefix)],capture_output=True,check=True,timeout=15)
            result=subprocess.run(['tesseract',str(prefix)+'.png','stdout','-l','spa+eng'],capture_output=True,text=True,check=True,timeout=20,env={**os.environ,'OMP_THREAD_LIMIT':'1'})
            paginas.append(result.stdout)
            if progreso:progreso('\n'.join(paginas))
        texto='\n'.join(paginas)
        if not texto.strip():raise ValueError('OCR sin texto reconocido')
        return texto


def excel(datos,url):
    """Tablas públicas: valores, rótulos y fechas; nunca ejecuta fórmulas ni macros."""
    if datos.startswith(b'PK'):
        with zipfile.ZipFile(BytesIO(datos)) as z:
            if sum(i.file_size for i in z.infolist())>32_000_000:raise ValueError('XLSX excede tamaño seguro')
        import openpyxl
        book=openpyxl.load_workbook(BytesIO(datos),read_only=True,data_only=True,keep_links=False)
        try:
            hojas=[];celdas=0
            for ws in book.worksheets:
                if (ws.max_row or 0)>100000 or (ws.max_column or 0)>1000:raise ValueError('Dimensiones XLSX exceden lectura segura')
                filas=[]
                for fila in ws.iter_rows(values_only=True):
                    celdas+=len(fila)
                    if celdas>500000:raise ValueError('Planilla supera presupuesto de celdas; lectura pendiente')
                    filas.append(fila)
                hojas.append((ws.title,filas))
        finally:book.close()
    else:
        import xlrd
        book=xlrd.open_workbook(file_contents=datos,on_demand=True)
        hojas=[]
        try:
            celdas=0
            for ws in book.sheets():
                celdas+=ws.nrows*ws.ncols
                if ws.nrows>100000 or celdas>500000:raise ValueError('XLS excede lectura segura')
                filas=[]
                for i in range(ws.nrows):
                    fila=[]
                    for c in ws.row(i):
                        if c.ctype==xlrd.XL_CELL_DATE:valor=xlrd.xldate.xldate_as_datetime(c.value,book.datemode)
                        else:valor=c.value
                        fila.append(valor)
                    filas.append(fila)
                hojas.append((ws.name,filas))
        finally:book.release_resources()
    from datetime import datetime,date
    def valor(v):
        if v is None:return ''
        if isinstance(v,(date,datetime)):return v.isoformat(sep=' ') if isinstance(v,datetime) else v.isoformat()
        if isinstance(v,float) and v.is_integer():return str(int(v))
        return str(v).replace('\n',' ').strip()
    texto=[]
    for nombre,filas in hojas:
        if len(filas)>100000:raise ValueError('Planilla supera límite de lectura segura; no se declara completa')
        texto.append(nombre)
        for fila in filas:
            celdas=[valor(c) for c in fila]
            if any(celdas):texto.append(' | '.join(celdas))
    if not texto:raise ValueError('Planilla sin celdas reconocibles')
    return '\n'.join(texto)


def leer(datos,tipo,url,ocr=False,progreso=None):
    if datos.startswith(b'%PDF-') or 'pdf' in tipo.lower():return pdf(datos,ocr,progreso)
    if tipo.lower().startswith(('image/jpeg','image/png')):
        if not ocr:raise ValueError('Imagen de publicación requiere OCR')
        with tempfile.TemporaryDirectory() as d:
            image=Path(d)/'publicacion';image.write_bytes(datos)
            r=subprocess.run(['tesseract',str(image),'stdout','-l','spa+eng'],capture_output=True,text=True,check=True,timeout=20,env={**os.environ,'OMP_THREAD_LIMIT':'1'})
            if not r.stdout.strip():raise ValueError('Imagen sin texto reconocido')
            return r.stdout
    if datos.startswith(b'\xd0\xcf\x11\xe0') or (datos.startswith(b'PK') and ('spreadsheet' in tipo.lower() or url.lower().split('?')[0].endswith('.xlsx'))):return excel(datos,url)
    if datos.startswith(b'PK'):
        with zipfile.ZipFile(BytesIO(datos)) as z:
            if 'xl/workbook.xml' in z.namelist():return excel(datos,url)
            if 'word/document.xml' in z.namelist():return docx(datos)
            if sum(i.file_size for i in z.infolist())>32_000_000:raise ValueError('ZIP excede tamaño seguro')
            textos=[]
            for i in z.infolist():
                nombre=i.filename.lower()
                if nombre.endswith(('.pdf','.docx','.txt','.xls','.xlsx')):
                    contenido=z.read(i)
                    textos.append(contenido.decode('utf-8') if nombre.endswith('.txt') else leer(contenido,'',nombre,ocr,progreso))
            if not textos:raise ValueError('ZIP sin documentos de contratación legibles')
            return '\n'.join(textos)
    raise ValueError('Formato de documento no admitido: '+tipo)


def ocr_selectivo(datos,max_paginas=4):
    """OCR por páginas y hash, con cupo explícito. Nunca cachea un parcial como completo."""
    from hashlib import sha256
    import json
    import cache_documentos as cache
    if not datos.startswith(b'%PDF-') or not 1<=max_paginas<=10:raise ValueError('OCR selectivo fuera de rango')
    digest=sha256(datos).hexdigest();base=cache.raiz()/'ocr-paginas-v1'/digest;base.mkdir(parents=True,exist_ok=True)
    native=subprocess.run(['pdftotext','-layout','-','-'],input=datos,capture_output=True,check=True,timeout=20).stdout.decode('utf-8').split('\f')
    with tempfile.TemporaryDirectory() as d:
        source=Path(d)/'documento.pdf';source.write_bytes(datos)
        info=subprocess.run(['pdfinfo',str(source)],capture_output=True,text=True,check=True,timeout=5)
        m=re.search(r'^Pages:\s*(\d+)',info.stdout,re.M)
        if not m:raise ValueError('OCR sin cantidad de páginas comprobable')
        total=int(m[1]);leidas={};necesarias=[];nativas=[]
        for n in range(1,total+1):
            t=native[n-1] if n<=len(native) else ''
            if len(re.sub(r'[^A-Za-zÁÉÍÓÚáéíóúÑñ]','',t))>=20:leidas[n]=t;nativas.append(n);continue
            p=base/(str(n)+'.json')
            if p.exists():
                r=json.loads(p.read_text())
                if r.get('sha256')==digest and r.get('pagina')==n:leidas[n]=r['texto'];continue
            necesarias.append(n)
        # Portada y últimas páginas primero, sin perder el número real de página.
        orden=necesarias[:1]+list(reversed(necesarias[1:]));procesadas=[]
        for n in orden[:max_paginas]:
            prefix=Path(d)/'pagina'
            subprocess.run(['pdftoppm','-f',str(n),'-l',str(n),'-scale-to','1600','-singlefile','-png',str(source),str(prefix)],capture_output=True,check=True,timeout=12)
            r=subprocess.run(['tesseract',str(prefix)+'.png','stdout','-l','spa+eng'],capture_output=True,text=True,check=True,timeout=15,env={**os.environ,'OMP_THREAD_LIMIT':'1'})
            leidas[n]=r.stdout;procesadas.append(n)
            cache.atomico(base/(str(n)+'.json'),{'sha256':digest,'pagina':n,'texto':r.stdout})
        meta={'sha256':digest,'total_paginas':total,'paginas_leidas':sorted(leidas),'paginas_nativas':nativas,'paginas_ocr_nuevas':procesadas,'paginas_pendientes':[n for n in range(1,total+1) if n not in leidas],'completa':len(leidas)==total}
        for n,t in leidas.items():cache.atomico(base/(str(n)+'.json'),{'sha256':digest,'pagina':n,'texto':t})
        cache.atomico(base/'indice.json',meta)
        return '\f'.join(leidas.get(n,'') for n in range(1,total+1)),meta


def ocr_en_cache(datos):
    """Reutiliza un OCR completo por hash sin consumir CPU ni cupo de nuevas páginas."""
    from hashlib import sha256
    import json
    import cache_documentos as cache
    digest=sha256(datos).hexdigest();base=cache.raiz()/'ocr-paginas-v1'/digest
    try:
        meta=json.loads((base/'indice.json').read_text())
        if not meta.get('completa') or meta.get('sha256')!=digest:return None
        paginas=[]
        for n in range(1,meta['total_paginas']+1):
            r=json.loads((base/(str(n)+'.json')).read_text())
            if r.get('sha256')!=digest or r.get('pagina')!=n:return None
            paginas.append(r['texto'])
        return '\f'.join(paginas),{**meta,'paginas_ocr_nuevas':[],'cache_completa':True}
    except (OSError,ValueError,KeyError,TypeError):return None
